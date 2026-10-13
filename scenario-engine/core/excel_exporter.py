"""
CWA Planning Automation - Excel Exporter
Complete corrected version for Option 3 Only
Includes full support for SequenceOrder + Relations sheets
"""
import pandas as pd
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from collections import defaultdict, deque
from pathlib import Path
import re
from config import ALT_POUR_ID_PROP
import math
import time as _time

from core.clash_position import classify_vertical_position


# ========================== HELPER FUNCTIONS ==========================

def _safe_int(value, default=999999):
    try:
        return int(str(value).strip())
    except Exception:
        return default


def _extract_level_sort_key(level_value, default=999999):
    if level_value is None:
        return default
    text = str(level_value).strip()
    if not text:
        return default
    match = re.search(r"[-+]?\d+", text)
    if not match:
        return default
    try:
        return int(match.group(0))
    except Exception:
        return default


def _logic_group_1_sort_rank(lg1_value, default=99):
    normalized = str(lg1_value or "").replace("_", " ").strip().upper()
    normalized = re.sub(r"\s*\(STAGE\s*2\)\s*$", "", normalized)
    priority = {
        "EXTERNAL WALL": 1, "INTERNAL WALL": 2, "SLAB": 3, "LANDING": 4,
        "COLUMN": 5, "BEAM": 6, "PARAPET": 7, "PEDESTAL": 8,
        "GULLIES": 9, "SCREED": 10, "UPSTAND": 11,
    }
    return priority.get(normalized, default)


def _building_zone_sort_rank(building_zone_value, default=99):
    normalized = str(building_zone_value or "").strip().upper()
    priority = {"DISCONNECTED": 1, "CONNECTED_HK": 2, "CONNECTED_HL": 3, "CONNECTED": 4}
    return priority.get(normalized, default)


def _logic_group_2_reserve_count(lg2_value, group_count):
    normalized = str(lg2_value or "").strip().upper()
    pct = 0.65 if "WALL" in normalized else 0.45
    return int(math.ceil(max(0, int(group_count or 0)) * pct))


def _id_to_int(value, default=999999):
    text = str(value or "").strip()
    if not text:
        return default
    match = re.match(r"^(\d+)", text)
    if not match:
        return default
    try:
        return int(match.group(1))
    except Exception:
        return default


def _build_reserved_placeholder_row(proposed_id_text):
    return {
        "No.": "",
        "PlanningProposedPourID": str(proposed_id_text or "").zfill(4),
        "PlanningProposedPourID_Override": "",
        "Entry Type": "RESERVED_PLACEHOLDER",
        "Placeholder Note": "Reserved number placeholder (available for future pours)",
        "Level": "", "Pour ID Number": "", ALT_POUR_ID_PROP: "", "FullPourID": "", "Pour Type": "",
        "General Pour Type": "", "Pour Volume m3": "", "Is External?": "",
        "Building Zone": "", "Concrete Stage": "",
        "Logic Group 1": "", "Logic Group 2": "", "Logic Group 3": "",
        "Max Z": "", "Min Z": "", "Include": "", "Building": "",
        "Centre of Gravity X": "", "Centre of Gravity Y": "", "Centre of Gravity Z": "",
        "GUID": "",
    }


def _normalize_proposed_id_text(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        try:
            int_value = int(value)
            return str(int_value).zfill(4) if int_value > 0 else ""
        except:
            return ""
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none"}:
        return ""
    if text.isdigit():
        return str(int(text)).zfill(4)
    if re.match(r"^\d+\.0+$", text):
        return str(int(float(text))).zfill(4)
    return ""


def _force_column_to_text(worksheet, column_names, header_row=1):
    """
    Force one or more columns to Text format (@) so leading zeros are preserved.
    """
    if isinstance(column_names, str):
        column_names = [column_names]

    for col_idx in range(1, worksheet.max_column + 1):
        header = str(worksheet.cell(row=header_row, column=col_idx).value or "").strip()
        if header in column_names:
            for row_idx in range(header_row + 1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.number_format = '@'
                if cell.value is not None:
                    val = str(cell.value).strip()
                    if val and val.lower() not in {"nan", "none"}:
                        cell.value = val
            return True
    return False


def _fast_column_widths(df):
    """Column width per header, computed from an in-memory DataFrame with
    vectorized pandas string ops — used to skip _apply_sheet_formatting_and_
    tables' slow fallback (re-reading every cell of every row back out of
    the just-saved .xlsx via openpyxl to find the longest value), which used
    to be the dominant cost of every Excel export (~11-16s on a ~3000-row
    workbook, measured) despite running unconditionally on every single run
    regardless of what data actually changed.

    Returns {header: max_char_length} — NOT yet clamped/scaled to a column
    width; _apply_sheet_formatting_and_tables applies the same
    "* 1.05 + 1, capped at 45" formula it always did, just against this
    precomputed length instead of a fresh cell-by-cell scan."""
    widths = {}
    if df is None:
        return widths
    for col in df.columns:
        header = str(col).strip()
        header_len = len(header)
        if df.empty or col not in df.columns:
            widths[header] = header_len
            continue
        # notna()-gated instead of a straight .astype(str) so NaN/None don't
        # count as the 3-4 char strings "nan"/"None" toward the max length.
        # astype(object) first: nullable dtypes (Int64...) reject "" as a fill.
        col_series = df[col].astype(object)
        str_lengths = col_series.where(col_series.notna(), "").astype(str).str.len()
        max_len = int(str_lengths.max()) if len(str_lengths) else 0
        widths[header] = max(header_len, max_len)
    return widths


def _apply_sheet_formatting_and_tables(wb, precomputed_widths=None):
    """Bolds/centers the header row, force-texts PlanningProposedPourID-style
    columns, sizes columns, and — as long as a sheet has at least one data
    row below its header — turns its range into a real Excel Table
    (displayName "tbl_<Sheet_Name>", spaces replaced with underscores),
    instead of leaving it a plain unstyled range. Applied to every sheet in
    whatever workbook is passed in; shared between the main per-model report
    and the Overrides workbook (SequenceOrder, SequenceOrder Template,
    SequenceOrder Conflicts, Numbering, Relations, PourIDOrder, ...) so both
    get the same filter arrows/structured references/banding.

    precomputed_widths: optional {sheet_name: {header: max_char_length}}
    (see _fast_column_widths) — sheets/columns present here skip the slow
    per-cell openpyxl scan entirely. Falls back to that scan for any sheet
    or column not covered (e.g. a caller that didn't pass one at all), so
    this stays correct even for callers that don't have a DataFrame handy."""
    precomputed_widths = precomputed_widths or {}
    FIXED_COLUMNS = {
        "No.": 5, "PlanningProposedPourID": 22, "PlanningProposedPourID_Override": 26,
        "Node Type": 14, "WBS Line ID": 22, "Parent ID": 22, "Activity ID": 20,
        "Activity Name": 58, "Original Duration": 16, "Phase Group": 16,
        "First Proposed Pour ID": 22, "Last Proposed Pour ID": 22, "Logical level 3": 45,
        "Total Pours": 12, "Full Pour ID": 22, "Days": 7, "Level": 9, "Ext": 6,
        ALT_POUR_ID_PROP: 16, f"A:{ALT_POUR_ID_PROP}": 16, f"B:{ALT_POUR_ID_PROP}": 16,
        "Logic Group 1": 25, "Logic Group 2": 35, "Logic Group 3": 35, "Type": 9,
        "Connected Pours": 11, "Include": 9, "TaskID_Override": 20, "WBS": 45,
        "Phase": 10, "Predecessors": 20, "Qty": 8, "Quantity Source": 16,
        "Stage": 10, "Zone": 7, "Min Z": 7, "Max Z": 7,
    }

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            cell.font = cell.font.copy(bold=True)

        _force_column_to_text(ws, ["PlanningProposedPourID", "PlanningProposedPourID_Override"])

        sheet_widths = precomputed_widths.get(sheet_name, {})
        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            header = str(ws.cell(row=1, column=col_idx).value or "").strip()
            if header in FIXED_COLUMNS:
                ws.column_dimensions[col_letter].width = FIXED_COLUMNS.get(header, 15)
            elif header in sheet_widths:
                max_length = sheet_widths[header]
                ws.column_dimensions[col_letter].width = min(max_length * 1.05 + 1, 45)
            else:
                max_length = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1))
                ws.column_dimensions[col_letter].width = min(max_length * 1.05 + 1, 45)

        if len(ws["A"]) > 1 or sheet_name == "SequenceOrder":
            # SequenceOrder is always turned into tbl_SequenceOrder even
            # when it's still just a header row (first-generation scaffold),
            # so the structured-reference table is there to paste into.
            table_range = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            tab = Table(displayName=f"tbl_{sheet_name.replace(' ', '_')}", ref=table_range)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                                showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            if sheet_name == "Activities":
                tab.showTotalsRow = True
            ws.add_table(tab)


# ========================== UPDATED load_overrides ==========================

def load_overrides(excel_path: str):
    overrides_path = Path(excel_path).parent / f"{Path(excel_path).stem.replace('_CWA_Planning', '')}_Overrides.xlsx"
    if not overrides_path.exists():
        return {}

    overrides = {}
    proposed_id_override_count = 0
    manual_order_count = 0

    try:
        # Overrides sheet
        df = pd.read_excel(str(overrides_path), sheet_name="Overrides", dtype=str)
        for _, row in df.fillna("").iterrows():
            guid = str(row.get("GUID", "") or "").strip()
            if not guid or guid == "MISSING_GUID":
                continue

            phase = str(row.get("Phase", "") or "").strip()
            task_id_override = str(row.get("TaskID_Override", "") or "").strip()
            task_name_override = str(row.get("TaskName_Override", "") or "").strip()

            proposed_id_override = _normalize_proposed_id_text(row.get("PlanningProposedPourID_Override", ""))
            if proposed_id_override:
                overrides.setdefault(guid, {})["PlanningProposedPourID"] = proposed_id_override
                proposed_id_override_count += 1

            if phase and (task_id_override or task_name_override):
                phase_key = f"{guid}_{phase}"
                overrides.setdefault(phase_key, {})
                if task_id_override:
                    overrides[phase_key]["TaskID"] = task_id_override
                if task_name_override:
                    overrides[phase_key]["TaskName"] = task_name_override

        # Numbering sheet
        try:
            df_numbering = pd.read_excel(str(overrides_path), sheet_name="Numbering", dtype=str)
            for _, row in df_numbering.fillna("").iterrows():
                guid = str(row.get("GUID", "") or "").strip()
                if not guid or guid == "MISSING_GUID":
                    continue
                proposed_id_override = _normalize_proposed_id_text(row.get("PlanningProposedPourID_Override", ""))
                if proposed_id_override:
                    overrides.setdefault(guid, {})["PlanningProposedPourID"] = proposed_id_override
                    proposed_id_override_count += 1
        except:
            pass

        # PourIDOrder sheet
        try:
            df_pour_order = pd.read_excel(str(overrides_path), sheet_name="PourIDOrder", dtype=str)
            if not df_pour_order.empty:
                manual_order = {}
                for idx, (_, values) in enumerate(df_pour_order.fillna("").iterrows(), 1):
                    guid = str(values.get("GUID", "") or "").strip()
                    if guid and guid != "MISSING_GUID" and guid not in manual_order:
                        manual_order[guid] = idx
                if manual_order:
                    overrides["__pour_order__"] = manual_order
                    manual_order_count = len(manual_order)
        except:
            pass

        # SequenceOrder sheet
        try:
            df_seq = pd.read_excel(str(overrides_path), sheet_name="SequenceOrder", dtype=str)
            seq_list = []
            for _, row in df_seq.iterrows():
                seq_list.append({
                    "Index": _safe_int(row.get("Index", 9999)),
                    "PourID": str(row.get("Pour ID", "")).strip(),
                    "GUID": str(row.get("GUID", "")).strip(),
                    "Phase": str(row.get("Phase", "") or row.get("Activity", "") or "").strip(),
                })
            overrides["__sequence_order__"] = sorted(seq_list, key=lambda x: x["Index"])
        except Exception as e:
            print(f"DEBUG: SequenceOrder load failed: {e}")
            overrides["__sequence_order__"] = []

        # Crews sheet (labour per phase / pour type, replacing config.CREW_TABLE rows)
        try:
            df_crews = pd.read_excel(str(overrides_path), sheet_name="Crews", dtype=str)
            overrides["__crews__"] = [
                {k: str(row.get(k, "") or "").strip() for k in ("Phase", "PourType", "Trade", "Crew")}
                for _, row in df_crews.fillna("").iterrows()
            ]
        except Exception:
            overrides["__crews__"] = []

        # Relations sheet
        try:
            df_rel = pd.read_excel(str(overrides_path), sheet_name="Relations", dtype=str)
            rel_list = []
            for _, row in df_rel.iterrows():
                a1 = str(row.get("ActivityID1", "")).strip()
                a2 = str(row.get("ActivityID2", "")).strip()
                if a1 and a2:
                    rel_list.append({
                        "ActivityID1": a1,
                        "ActivityID2": a2,
                        "Relation": str(row.get("Relation", "FS")).strip().upper(),
                        "Lag": float(row.get("Lag", 0) or 0)
                    })
            overrides["__relations__"] = rel_list
        except:
            overrides["__relations__"] = []

    except Exception as exc:
        print(f"Warning: could not fully read overrides file {overrides_path.name}: {exc}")
        return overrides

    print(f"Loaded overrides: {overrides_path.name} "
          f"(Proposed ID overrides: {proposed_id_override_count}, Manual sequence: {manual_order_count})")
    return overrides


# ========================== NUMBERING & SUMMARY FUNCTIONS ==========================

def _build_numbering_summary_from_records(records, numbering_mode="sequential"):
    groups = defaultdict(list)
    group_meta = {}

    for rec in records or []:
        if numbering_mode == "stage_priority":
            summary_key = (
                str(rec.get("conc_stage") or "").strip().upper(),
                str(rec.get("general_pour_type") or "").strip().upper(),
                str(rec.get("pour_type") or "").strip().upper(),
                str(rec.get("level_code") or rec.get("level_display") or "").strip().upper(),
                "YES" if bool(rec.get("is_external", False)) else "NO",
                str(rec.get("building_zone") or "").strip().upper(),
                str(rec.get("lg1") or "").strip(),
            )
        else:
            summary_key = str(rec.get("lg3") or rec.get("lg2") or "").strip()

        proposed_id = str(rec.get("planning_proposed_pour_id") or "").strip()
        if not summary_key or not proposed_id:
            continue
        groups[summary_key].append(proposed_id.zfill(4))
        group_meta.setdefault(summary_key, rec)

    has_spare = (numbering_mode == "stage_priority") or any("_sr_spare" in r for r in (records or []))

    summary_rows = []
    for summary_key, ids in groups.items():
        ordered_ids = sorted(ids, key=lambda x: _safe_int(x))
        if numbering_mode == "stage_priority":
            meta = group_meta.get(summary_key, {})
            row = {
                "Concrete Stage": summary_key[0],
                "General Pour Type": summary_key[1],
                "Pour Type": summary_key[2],
                "Level": summary_key[3],
                "Is External?": summary_key[4],
                "Building Zone": summary_key[5],
                "Logic Group 1": summary_key[6],
                "Logic Group 2": str(meta.get("lg2") or "").strip(),
                "First Proposed Pour ID": ordered_ids[0],
                "Last Proposed Pour ID": ordered_ids[-1],
                "Total Pours": len(ids),
            }
        else:
            meta = group_meta.get(summary_key, {})
            row = {
                "Logical level 3": summary_key,
                "Logic Group 1": str(meta.get("lg1") or "").strip(),
                "Logic Group 2": str(meta.get("lg2") or "").strip(),
                "First Proposed Pour ID": ordered_ids[0],
                "Last Proposed Pour ID": ordered_ids[-1],
                "Total Pours": len(ids),
            }

        if numbering_mode == "stage_priority":
            row["Spare IDs"] = 0
        else:
            row["Spare IDs"] = meta.get("_sr_spare", "")
        if has_spare:
            row["Next Level Starts At"] = meta.get("_sr_next_start", "")
        summary_rows.append(row)

    if numbering_mode == "stage_priority":
        base_cols = ["Concrete Stage", "General Pour Type", "Pour Type", "Level", "Is External?",
                     "Building Zone", "Logic Group 1", "Logic Group 2", "First Proposed Pour ID",
                     "Last Proposed Pour ID", "Total Pours", "Spare IDs"]
    else:
        base_cols = ["Logical level 3", "Logic Group 1", "Logic Group 2", "First Proposed Pour ID",
                     "Last Proposed Pour ID", "Total Pours", "Spare IDs"]
    extra_cols = ["Next Level Starts At"] if has_spare else []

    if not summary_rows:
        return pd.DataFrame(columns=base_cols + extra_cols)

    df_summary = pd.DataFrame(summary_rows)

    if numbering_mode == "stage_priority":
        df_summary["_LevelSortKey"] = df_summary["Level"].apply(_extract_level_sort_key)
        df_summary["_LG1SortKey"] = df_summary["Logic Group 1"].apply(_logic_group_1_sort_rank)
        df_summary["_ZoneSortKey"] = df_summary["Building Zone"].apply(_building_zone_sort_rank)
        sort_cols = ["First Proposed Pour ID", "_LevelSortKey", "_LG1SortKey", "_ZoneSortKey",
                     "Concrete Stage", "General Pour Type", "Pour Type", "Is External?",
                     "Building Zone", "Logic Group 1", "Logic Group 2"]
        df_summary = df_summary.sort_values(by=sort_cols, ascending=[True] * len(sort_cols)).reset_index(drop=True)
        first_ids = df_summary["First Proposed Pour ID"].tolist()
        last_ids = df_summary["Last Proposed Pour ID"].tolist()
        actual_spares = []
        for idx, last_id in enumerate(last_ids):
            if idx >= len(first_ids) - 1:
                actual_spares.append(0)
                continue
            current_last = _id_to_int(last_id)
            next_first = _id_to_int(first_ids[idx + 1])
            actual_spares.append(max(0, next_first - current_last - 1) if current_last < 999999 and next_first < 999999 else 0)
        df_summary["Spare IDs"] = actual_spares
        return df_summary.drop(columns=["_LevelSortKey", "_LG1SortKey", "_ZoneSortKey"], errors="ignore")

    df_summary["_ZoneSortKey"] = df_summary["Logic Group 2"].apply(_building_zone_sort_rank)
    sort_cols = ["First Proposed Pour ID", "_ZoneSortKey", "Logical level 3", "Logic Group 1", "Logic Group 2"]
    return df_summary.sort_values(by=sort_cols, ascending=[True] * len(sort_cols)).reset_index(drop=True).drop(columns=["_ZoneSortKey"], errors="ignore")


def _build_numbering_summary_from_activities(df_activities):
    if df_activities is None or df_activities.empty:
        return pd.DataFrame()

    df = df_activities.copy()
    if "PlanningProposedPourID" not in df.columns:
        return pd.DataFrame()

    df["PlanningProposedPourID"] = df["PlanningProposedPourID"].apply(_normalize_proposed_id_text)
    df = df[df["PlanningProposedPourID"].astype(str).str.strip() != ""]
    if df.empty:
        return pd.DataFrame()

    stage_cols = [
        "Stage", "Type", "Pour Type", "Level", "Ext", "Logic Group 2", "Logic Group 1"
    ]
    has_stage_priority_columns = all(c in df.columns for c in stage_cols)

    if has_stage_priority_columns:
        grouped = (
            df.groupby(stage_cols, dropna=False)["PlanningProposedPourID"]
            .apply(list)
            .reset_index(name="_ids")
        )

        rows = []
        for _, row in grouped.iterrows():
            ordered_ids = sorted([str(v).zfill(4) for v in row["_ids"]], key=lambda x: _safe_int(x))
            rows.append({
                "Concrete Stage": str(row["Stage"] or "").strip().upper(),
                "General Pour Type": str(row["Type"] or "").strip().upper(),
                "Pour Type": str(row["Pour Type"] or "").strip().upper(),
                "Level": str(row["Level"] or "").strip().upper(),
                "Is External?": "YES" if str(row["Ext"] or "").strip().upper() == "YES" else "NO",
                "Building Zone": str(row["Logic Group 2"] or "").strip().upper(),
                "Logic Group 1": str(row["Logic Group 1"] or "").strip(),
                "Logic Group 2": str(row["Logic Group 2"] or "").strip(),
                "First Proposed Pour ID": ordered_ids[0],
                "Last Proposed Pour ID": ordered_ids[-1],
                "Total Pours": len(ordered_ids),
            })

        if not rows:
            return pd.DataFrame(columns=[
                "Concrete Stage", "General Pour Type", "Pour Type", "Level", "Is External?",
                "Building Zone", "Logic Group 1", "Logic Group 2", "First Proposed Pour ID",
                "Last Proposed Pour ID", "Total Pours", "Spare IDs"
            ])

        df_summary = pd.DataFrame(rows)
        df_summary["_LevelSortKey"] = df_summary["Level"].apply(_extract_level_sort_key)
        df_summary["_LG1SortKey"] = df_summary["Logic Group 1"].apply(_logic_group_1_sort_rank)
        df_summary["_ZoneSortKey"] = df_summary["Building Zone"].apply(_building_zone_sort_rank)
        sort_cols = [
            "First Proposed Pour ID", "_LevelSortKey", "_LG1SortKey", "_ZoneSortKey",
            "Concrete Stage", "General Pour Type", "Pour Type", "Is External?",
            "Building Zone", "Logic Group 1", "Logic Group 2"
        ]
        df_summary = df_summary.sort_values(by=sort_cols, ascending=[True] * len(sort_cols)).reset_index(drop=True)

        first_ids = df_summary["First Proposed Pour ID"].tolist()
        last_ids = df_summary["Last Proposed Pour ID"].tolist()
        actual_spares = []
        for idx, last_id in enumerate(last_ids):
            if idx >= len(first_ids) - 1:
                actual_spares.append(0)
                continue
            current_last = _id_to_int(last_id)
            next_first = _id_to_int(first_ids[idx + 1])
            actual_spares.append(max(0, next_first - current_last - 1) if current_last < 999999 and next_first < 999999 else 0)
        df_summary["Spare IDs"] = actual_spares
        return df_summary.drop(columns=["_LevelSortKey", "_LG1SortKey", "_ZoneSortKey"], errors="ignore")

    df_fallback = df.copy()
    if "Full Pour ID" in df_fallback.columns:
        df_fallback["Logical level 3"] = df_fallback["Full Pour ID"].astype(str).str.strip()
    else:
        df_fallback["Logical level 3"] = ""

    if "Logic Group 1" not in df_fallback.columns:
        df_fallback["Logic Group 1"] = ""
    if "Logic Group 2" not in df_fallback.columns:
        df_fallback["Logic Group 2"] = ""

    grouped = (
        df_fallback.groupby(["Logical level 3", "Logic Group 1", "Logic Group 2"], dropna=False)["PlanningProposedPourID"]
        .apply(list)
        .reset_index(name="_ids")
    )

    rows = []
    for _, row in grouped.iterrows():
        logical_level_3 = str(row["Logical level 3"] or "").strip()
        if not logical_level_3:
            continue
        ordered_ids = sorted([str(v).zfill(4) for v in row["_ids"]], key=lambda x: _safe_int(x))
        rows.append({
            "Logical level 3": logical_level_3,
            "Logic Group 1": str(row["Logic Group 1"] or "").strip(),
            "Logic Group 2": str(row["Logic Group 2"] or "").strip(),
            "First Proposed Pour ID": ordered_ids[0],
            "Last Proposed Pour ID": ordered_ids[-1],
            "Total Pours": len(ordered_ids),
            "Spare IDs": "",
        })

    if not rows:
        return pd.DataFrame(columns=[
            "Logical level 3", "Logic Group 1", "Logic Group 2", "First Proposed Pour ID",
            "Last Proposed Pour ID", "Total Pours", "Spare IDs"
        ])

    df_summary = pd.DataFrame(rows)
    df_summary["_ZoneSortKey"] = df_summary["Logic Group 2"].apply(_building_zone_sort_rank)
    sort_cols = ["First Proposed Pour ID", "_ZoneSortKey", "Logical level 3", "Logic Group 1", "Logic Group 2"]
    return df_summary.sort_values(by=sort_cols, ascending=[True] * len(sort_cols)).reset_index(drop=True).drop(columns=["_ZoneSortKey"], errors="ignore")


def _build_numbering_table_from_records(records, overrides=None, include_override_column=False, numbering_mode="sequential"):
    if overrides is None:
        overrides = {}

    ordered_columns = [
        "No.", "PlanningProposedPourID", "PlanningProposedPourID_Override", "Entry Type",
        "Placeholder Note", "Level", "Pour ID Number", ALT_POUR_ID_PROP, "FullPourID", "Pour Type",
        "General Pour Type", "Pour Volume m3", "Is External?", "Building Zone",
        "Concrete Stage", "Logic Group 1", "Logic Group 2", "Logic Group 3", "Max Z", "Min Z",
        "Include", "Building", "Centre of Gravity X", "Centre of Gravity Y", "Centre of Gravity Z", "GUID"
    ]

    if not records:
        return pd.DataFrame(columns=ordered_columns)

    rows = []
    for rec in records:
        element = rec.get("element")
        guid = getattr(element, "GlobalId", "MISSING_GUID") if element else "MISSING_GUID"
        props = rec.get("cwa_pour_props", {}) or {}

        proposed_id = _normalize_proposed_id_text(rec.get("planning_proposed_pour_id", ""))
        proposed_id_override = _normalize_proposed_id_text(
            overrides.get(guid, {}).get("PlanningProposedPourID", "")
        ) if include_override_column else ""

        level_value = props.get("Level", rec.get("level_display", ""))
        pour_id_number = props.get("Pour ID Number", rec.get("pour_id", ""))
        alt_pour_id = rec.get("alt_pour_id", rec.get("original_pour_id", ""))
        full_pour_id = rec.get("full_pour_id", "") or props.get("Full Pour ID", "") or props.get("Full Pour ID String", "")
        pour_type = props.get("Pour Type", rec.get("pour_type", ""))
        general_pour_type = props.get("General Pour Type", "")
        pour_volume = props.get("Pour Volume m3", rec.get("volume", ""))
        is_external = props.get("Is External?", "YES" if rec.get("is_external", False) else "NO")
        building_zone = props.get("Building Zone", rec.get("building_zone", ""))
        concrete_stage = props.get("Concrete Stage", rec.get("conc_stage", ""))
        cog_x = props.get("Centre of Gravity X", rec.get("coord_x", ""))
        cog_y = props.get("Centre of Gravity Y", rec.get("coord_y", ""))
        cog_z = props.get("Centre of Gravity Z", rec.get("coord_z", ""))

        row = {
            "No.": "",
            "PlanningProposedPourID": proposed_id,
            "PlanningProposedPourID_Override": proposed_id_override,
            "Entry Type": "POUR",
            "Placeholder Note": "",
            "_Option3OrderRank": rec.get("_option3_order_rank", 10**6),
            "Level": level_value,
            "Pour ID Number": pour_id_number,
            ALT_POUR_ID_PROP: alt_pour_id,
            "FullPourID": full_pour_id,
            "Pour Type": pour_type,
            "General Pour Type": general_pour_type,
            "Pour Volume m3": pour_volume,
            "Is External?": is_external,
            "Building Zone": building_zone,
            "Concrete Stage": concrete_stage,
            "GUID": guid,
            "Logic Group 1": rec.get("lg1", ""),
            "Logic Group 2": rec.get("lg2", ""),
            "Logic Group 3": rec.get("lg3", ""),
            "Max Z": rec.get("max_elevation", ""),
            "Min Z": rec.get("min_elevation", ""),
            "Include": rec.get("include", "NO"),
            "Building": rec.get("building", ""),
            "Centre of Gravity X": cog_x,
            "Centre of Gravity Y": cog_y,
            "Centre of Gravity Z": cog_z,
        }
        rows.append(row)

    df = pd.DataFrame(rows)

    if "PlanningProposedPourID" in df.columns:
        df["PlanningProposedPourID"] = df["PlanningProposedPourID"].apply(_normalize_proposed_id_text)
    if "PlanningProposedPourID_Override" in df.columns:
        df["PlanningProposedPourID_Override"] = df["PlanningProposedPourID_Override"].apply(_normalize_proposed_id_text)

    if numbering_mode == "stage_priority":
        df = df.sort_values(
            by=["Include", "Building", "Level", "_Option3OrderRank", "PlanningProposedPourID"],
            ascending=[False, True, True, True, True]
        ).reset_index(drop=True)

        used_ids = sorted({_id_to_int(v) for v in df["PlanningProposedPourID"].tolist() if _id_to_int(v) < 999999})
        if used_ids:
            min_used = used_ids[0]
            max_used = used_ids[-1]
            used_set = set(used_ids)
            missing_ids = [i for i in range(min_used, max_used + 1) if i not in used_set]
            if missing_ids:
                placeholder_rows = [_build_reserved_placeholder_row(str(mid).zfill(4)) for mid in missing_ids]
                df = pd.concat([df, pd.DataFrame(placeholder_rows)], ignore_index=True)

        df["_IdSort"] = df["PlanningProposedPourID"].apply(_id_to_int)
        df["_EntryTypeSort"] = df["Entry Type"].apply(lambda v: 0 if str(v) == "POUR" else 1)
        df = df.sort_values(
            by=["_IdSort", "_EntryTypeSort", "Building", "Level", "_Option3OrderRank", "GUID"],
            ascending=[True, True, True, True, True, True]
        ).reset_index(drop=True)
        df = df.drop(columns=["_IdSort", "_EntryTypeSort"], errors="ignore")
    else:
        sort_cols = ["GUID", "PlanningProposedPourID"]
        if "PlanningProposedPourID_Override" in df.columns:
            sort_cols = ["PlanningProposedPourID_Override", "GUID", "PlanningProposedPourID"]
        df = df.sort_values(by=sort_cols, ascending=True).reset_index(drop=True)

    df = df[[c for c in ordered_columns if c in df.columns]]
    df["No."] = range(1, len(df) + 1)
    return df


def _build_clash_relations_from_records(clash_results, records):
    if clash_results is None:
        clash_results = []
    if records is None:
        records = []

    pour_lookup = {r["element"].GlobalId: r for r in records if "element" in r and hasattr(r["element"], "GlobalId")}
    seen_pairs = set()
    rows = []

    for c in clash_results:
        a_guid = c.get("Pour A GUID") or c.get("a_global_id", "")
        b_guid = c.get("Pour B GUID") or c.get("b_global_id", "")
        if not a_guid or not b_guid:
            continue
        pair_key = tuple(sorted((a_guid, b_guid)))
        if pair_key in seen_pairs:
            continue
        seen_pairs.add(pair_key)

        a = pour_lookup.get(a_guid, {})
        b = pour_lookup.get(b_guid, {})

        a_type = str(a.get("pour_type", "") or "").replace(" ", "_").upper()
        b_type = str(b.get("pour_type", "") or "").replace(" ", "_").upper()

        same_group = (str(a.get("building", "")) == str(b.get("building", "")) and
                      str(a.get("level_code", "")) == str(b.get("level_code", "")) and
                      str(a.get("lg3", "")) == str(b.get("lg3", "")))

        both_wall_or_slab = (("WALL" in a_type or "SLAB" in a_type) and
                             ("WALL" in b_type or "SLAB" in b_type))

        rows.append({
            "No.": len(rows) + 1,
            "A GUID": a_guid,
            "A Proposed ID": str(a.get("planning_proposed_pour_id", "") or "").zfill(4),
            "A Building": a.get("building", ""),
            "A Level": a.get("level_code", ""),
            "A Group": a.get("lg3", ""),
            "A Type": a.get("pour_type", ""),
            "A Pour ID": a.get("pour_id", ""),
            f"A {ALT_POUR_ID_PROP}": a.get("alt_pour_id", a.get("original_pour_id", "")),
            "A Min Z": a.get("min_elevation", ""),
            "A Max Z": a.get("max_elevation", ""),
            "B GUID": b_guid,
            "B Proposed ID": str(b.get("planning_proposed_pour_id", "") or "").zfill(4),
            "B Building": b.get("building", ""),
            "B Level": b.get("level_code", ""),
            "B Group": b.get("lg3", ""),
            "B Type": b.get("pour_type", ""),
            "B Pour ID": b.get("pour_id", ""),
            f"B {ALT_POUR_ID_PROP}": b.get("alt_pour_id", b.get("original_pour_id", "")),
            "B Min Z": b.get("min_elevation", ""),
            "B Max Z": b.get("max_elevation", ""),
            "Position (B rel. A)": classify_vertical_position(a, b) if a and b else "Undetermined",
            "Same Group": "YES" if same_group else "NO",
            "Chain Candidate": "YES" if same_group and both_wall_or_slab else "NO",
            "Distance (mm)": c.get("Distance (mm)", ""),
            "Distance (m)": c.get("Distance (m)", ""),
            "Violation": c.get("Violation", "< 5mm clearance"),
        })

    if not rows:
        return pd.DataFrame(columns=["No.", "A GUID", "A Proposed ID", "A Building", "A Level", "A Group", "A Type",
                                     "A Pour ID", f"A {ALT_POUR_ID_PROP}", "A Min Z", "A Max Z", "B GUID", "B Proposed ID", "B Building", "B Level",
                                     "B Group", "B Type", "B Pour ID", f"B {ALT_POUR_ID_PROP}", "B Min Z", "B Max Z", "Position (B rel. A)",
                                     "Same Group", "Chain Candidate", "Distance (mm)", "Distance (m)", "Violation"])

    df_relations = pd.DataFrame(rows)
    df_relations = df_relations.sort_values(
        by=["Chain Candidate", "Same Group", "A Building", "A Level", "A Group", "A Proposed ID", "B Proposed ID"],
        ascending=[False, False, True, True, True, True, True]
    ).reset_index(drop=True)
    df_relations["No."] = range(1, len(df_relations) + 1)
    return df_relations


def _build_duplicate_pour_ids_table(records):
    if not records:
        return pd.DataFrame()

    from collections import defaultdict as _dd
    groups = _dd(list)
    for rec in records:
        fid = str(rec.get("original_pour_id") or rec.get("pour_id") or "").strip()
        if fid:
            groups[fid].append(rec)

    if not any(len(g) >= 2 for g in groups.values()):
        return pd.DataFrame(columns=["No.", "Former Pour ID", "Objects with same ID", "GUID",
                                     "Proposed Pour ID", "Building", "Level", "Pour Type"])

    rows = []
    for fid, group in groups.items():
        if len(group) < 2:
            continue
        for rec in group:
            element = rec.get("element")
            guid = getattr(element, "GlobalId", "MISSING_GUID") if element else "MISSING_GUID"
            rows.append({
                "No.": "",
                "Former Pour ID": fid,
                "Objects with same ID": len(group),
                "GUID": guid,
                "Proposed Pour ID": str(rec.get("planning_proposed_pour_id") or rec.get("pour_id") or "").strip(),
                "Building": rec.get("building", ""),
                "Level": rec.get("level_display", ""),
                "Pour Type": rec.get("pour_type", ""),
            })

    if not rows:
        return pd.DataFrame(columns=["No.", "Former Pour ID", "Objects with same ID", "GUID",
                                     "Proposed Pour ID", "Building", "Level", "Pour Type"])

    df = pd.DataFrame(rows)
    df["No."] = range(1, len(df) + 1)
    return df


def _build_building_mapping_table(records):
    """
    Audit table: for each distinct original IfcBuilding name found in the
    source model, show the final building code it resolves to (via
    config.BUILDING_NAME_MAP or the Building/Unit property fallback), the
    element count, and whether it was merged with other original names that
    resolved to the same final code.
    """
    if not records:
        return pd.DataFrame()

    from collections import defaultdict as _dd
    groups = _dd(lambda: {"count": 0, "mapped": False})
    for rec in records:
        original = str(rec.get("original_building_name") or "").strip() or "(none / not found)"
        final = str(rec.get("building") or "").strip()
        key = (original, final)
        groups[key]["count"] += 1
        groups[key]["mapped"] = groups[key]["mapped"] or bool(rec.get("building_mapped"))

    final_to_originals = _dd(set)
    for (original, final) in groups.keys():
        final_to_originals[final].add(original)

    rows = []
    for (original, final), info in sorted(groups.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        others = sorted(o for o in final_to_originals[final] if o != original)
        rows.append({
            "No.": "",
            "Original IfcBuilding Name": original,
            "Final Building Code": final,
            "Source": "BUILDING_NAME_MAP" if info["mapped"] else "Building/Unit property",
            "Element Count": info["count"],
            "Merged With": ", ".join(others) if others else "",
        })

    df = pd.DataFrame(rows)
    if not df.empty:
        df["No."] = range(1, len(df) + 1)
    return df


def _compute_task_depths(task_predecessors, all_task_ids):
    """Longest-path depth (0 = no predecessors) for every task_id in
    all_task_ids, from the merged predecessor graph (task_id -> [predecessor
    task_id, ...], as built in main.py from intra-pour + clash-derived
    edges). Tasks that only reach back into a dependency cycle are left at
    whatever depth Kahn's algorithm got to before stalling — a safe
    fallback (they just won't get pushed later than their reachable
    predecessors), rather than raising on a graph that isn't a strict DAG."""
    task_predecessors = task_predecessors or {}
    successors = defaultdict(list)
    in_degree = {t: 0 for t in all_task_ids}
    for t in all_task_ids:
        for p in dict.fromkeys(task_predecessors.get(t, [])):
            if p in in_degree:
                in_degree[t] += 1
                successors[p].append(t)

    depth = {t: 0 for t in all_task_ids}
    remaining = dict(in_degree)
    queue = deque([t for t in all_task_ids if in_degree[t] == 0])
    while queue:
        t = queue.popleft()
        for s in successors.get(t, []):
            if depth[t] + 1 > depth[s]:
                depth[s] = depth[t] + 1
            remaining[s] -= 1
            if remaining[s] == 0:
                queue.append(s)

    return depth


def _representative_pour_task(task_map):
    """The one task per pour that "SequenceOrder"'s Index should actually
    track — the concrete-pour event itself (or, for Heavy Precast/
    Removable/Stair, the Installation "drop" that stands in for it — per
    the user's own framing: "pour activities paired with precast
    installation"). Phase-level detail (Rebar/Formwork/Stitching/Curing/
    Propping...) is already fully automatic via intra-pour + clash-derived
    predecessors and isn't what a strategic override like "all external
    walls + precast up to level N before inner structures" is about."""
    return (
        task_map.get("Concrete")
        or task_map.get("Installation")
        or task_map.get("Rebar, Formwork and Concrete")
        or task_map.get("Stitching")
        or task_map.get("Rebar")
        or next(iter(task_map.values()), None)
    )


def _guid_to_representative_task_map(activities):
    guid_to_tasks = defaultdict(dict)
    for act in activities or []:
        task_id = act.get("task_id")
        if not task_id:
            continue
        rec = act.get("rec", {})
        element = rec.get("element")
        guid = str(getattr(element, "GlobalId", "") or "").strip() if element else ""
        phase = str(act.get("phase", "") or "").strip()
        if guid and phase:
            guid_to_tasks[guid][phase] = task_id
    return guid_to_tasks


def _compute_pour_sequence_layers(activities, task_predecessors, task_start_dates=None):
    """Maps each pour GUID to a dense-ranked "layer" number (1, 2, 3, ...)
    reflecting the order pours actually happen in, keyed off each pour's own
    representative pour-task (see _representative_pour_task), so
    "SequenceOrder Template" reads as a sequence of concrete pours (Heavy
    Precast standing in via its Installation) rather than a phase-level
    chain. Pours landing in the same bucket share a layer (start together —
    matching "SequenceOrder"'s own same-Index == Start-to-Start semantics).

    Prefers task_start_dates ({task_id: ISO datetime string}) — the REAL
    CPM-computed PlannedStartDate already sitting in the XML after
    apply_cpm_to_xml runs — as the ranking key, bucketing pours by that
    task's actual calendar date. This is what makes the template match the
    real schedule: ranking by raw predecessor-graph hop count instead (the
    previous approach) ignores duration entirely, so a chain with one
    28-day Curing hop (Heavy Precast) looks just as "close" as a chain of
    several 1-day hops (a plain wall) — badly scrambling unrelated
    buildings/levels relative to each other despite the real schedule
    placing them nowhere near the same point in time.

    Falls back to unweighted predecessor-graph depth only when no dates are
    available (e.g. CPM was skipped) — still better than an arbitrary sort,
    just not duration-aware.

    Returns {} if activities/task_predecessors aren't available, so callers
    can fall back to the old plain-sort behavior."""
    if not activities or task_predecessors is None:
        return {}

    guid_to_tasks = _guid_to_representative_task_map(activities)
    if not guid_to_tasks:
        return {}

    if task_start_dates:
        guid_to_key = {}
        for guid, task_map in guid_to_tasks.items():
            representative_task = _representative_pour_task(task_map)
            start_date = task_start_dates.get(representative_task) if representative_task else None
            if start_date:
                guid_to_key[guid] = start_date
    else:
        all_task_ids = {act.get("task_id") for act in activities if act.get("task_id")}
        depths = _compute_task_depths(task_predecessors, all_task_ids)
        guid_to_key = {}
        for guid, task_map in guid_to_tasks.items():
            representative_task = _representative_pour_task(task_map)
            if representative_task:
                guid_to_key[guid] = depths.get(representative_task, 0)

    distinct_keys = sorted(set(guid_to_key.values()))
    layer_rank = {raw_key: rank + 1 for rank, raw_key in enumerate(distinct_keys)}
    return {guid: layer_rank[raw_key] for guid, raw_key in guid_to_key.items()}


def _build_sequence_order_template(records, activities=None, task_predecessors=None, task_start_dates=None):
    ordered_columns = [
        "Index", "Pour ID", "Level", "General Level 3", "Pour Type", "Concrete Stage", "Pour Volume m3", "Height",
        "Logic Group 1", "Logic Group 2", "Logic Group 3",
        ALT_POUR_ID_PROP, "PlanningProposedPourID", "GUID",
    ]

    if not records:
        return pd.DataFrame(columns=ordered_columns)

    guid_to_layer = _compute_pour_sequence_layers(activities, task_predecessors, task_start_dates)

    rows = []
    for rec in records:
        element = rec.get("element")
        guid = getattr(element, "GlobalId", "MISSING_GUID") if element else "MISSING_GUID"
        proposed_id = _normalize_proposed_id_text(rec.get("planning_proposed_pour_id", ""))
        pour_id = str(rec.get("pour_id", "") or "").strip()
        level_value = str(rec.get("level_code", rec.get("level_display", "")) or "").strip()

        rows.append({
            "Index": guid_to_layer.get(guid, None),
            "Pour ID": pour_id,
            ALT_POUR_ID_PROP: str(rec.get("alt_pour_id", rec.get("original_pour_id", "")) or "").strip(),
            "GUID": guid,
            "PlanningProposedPourID": proposed_id,
            "Level": level_value,
            "General Level 3": str(rec.get("general_level_3", "") or "").strip(),
            "Pour Type": str(rec.get("pour_type", "") or "").strip(),
            "Concrete Stage": str(rec.get("conc_stage", "") or "").strip(),
            "Pour Volume m3": rec.get("volume", ""),
            "Height": rec.get("bbox_height", ""),
            "Logic Group 1": str(rec.get("lg1", "") or "").strip(),
            "Logic Group 2": str(rec.get("lg2", "") or "").strip(),
            "Logic Group 3": str(rec.get("lg3", "") or "").strip(),
            "_IncludeSortKey": 0 if str(rec.get("include", "NO") or "NO").strip().upper() == "YES" else 1,
            "_LayerSortKey": guid_to_layer.get(guid, 999999),
            "_BuildingSortKey": str(rec.get("building", "") or "").strip(),
            "_LevelSortKey": _extract_level_sort_key(level_value),
            "_IdSortKey": _id_to_int(proposed_id),
        })

    df = pd.DataFrame(rows)
    df = df.sort_values(
        by=["_IncludeSortKey", "_LayerSortKey", "_BuildingSortKey", "_LevelSortKey", "_IdSortKey", "GUID"],
        ascending=[True, True, True, True, True, True]
    ).reset_index(drop=True)

    if not guid_to_layer:
        # No computed graph available (e.g. called without activities/
        # task_predecessors) — fall back to the old plain sequential
        # numbering rather than leaving Index blank for every row.
        df["Index"] = range(1, len(df) + 1)

    return _coerce_index_column(df[[c for c in ordered_columns if c in df.columns]])


_AUTOMATIC_SOURCE_LABELS = {
    "slab_support": "Clash-derived slab support",
    "level_gate": "Level gate",
    "heavy_precast_wall": "Heavy Precast <-> wall",
    "intra": "Intra-pour phase chain",
}


def _build_sequence_order_conflicts_table(activity_list, skipped_edges_by_source):
    """One row per link dropped to prevent a logic cycle, keyed by which
    source lost. SequenceOrder (the user's explicit, manual override) is
    emitted before every automatic source (level-gate, clash-derived slab
    support, Heavy Precast <-> wall) in write_activities_and_relationships,
    so a dropped automatic-source link here means SequenceOrder won a real
    conflict — expected behavior, not a bug, whenever the user's manual
    priority genuinely contradicts what the automatic rule would have
    required. A dropped "sequence" link (rare) means two SequenceOrder rows
    directly contradict each other, or contradict a pour's own intra-pour
    phase chain — that one IS worth the user's attention, since neither side
    of it made it into the schedule."""
    ordered_columns = [
        "No.", "Source", "Winner", "Requested Order", "Pred Task ID", "Pred Pour ID", "Pred Building",
        "Pred Level", "Pred Phase", "Succ Task ID", "Succ Pour ID", "Succ Building", "Succ Level", "Succ Phase",
    ]

    skipped_edges_by_source = skipped_edges_by_source or {}
    if not any(skipped_edges_by_source.values()):
        return pd.DataFrame(columns=ordered_columns)

    task_info = {}
    for act in activity_list or []:
        task_id = act.get("task_id")
        if not task_id:
            continue
        rec = act.get("rec", {})
        task_info[task_id] = {
            "Pour ID": rec.get("pour_id", ""),
            "Building": rec.get("building", ""),
            "Level": rec.get("level_code", rec.get("level_display", "")),
            "Phase": act.get("phase", ""),
        }

    rows = []
    for source, edges in skipped_edges_by_source.items():
        source_label = _AUTOMATIC_SOURCE_LABELS.get(source, source)
        winner = "Another SequenceOrder row" if source == "sequence" else "SequenceOrder"
        for pred_id, succ_id in edges:
            pred_info = task_info.get(pred_id, {})
            succ_info = task_info.get(succ_id, {})
            rows.append({
                "No.": len(rows) + 1,
                "Source": source_label,
                "Winner": winner,
                "Requested Order": f"{pred_id} before {succ_id}",
                "Pred Task ID": pred_id,
                "Pred Pour ID": pred_info.get("Pour ID", ""),
                "Pred Building": pred_info.get("Building", ""),
                "Pred Level": pred_info.get("Level", ""),
                "Pred Phase": pred_info.get("Phase", ""),
                "Succ Task ID": succ_id,
                "Succ Pour ID": succ_info.get("Pour ID", ""),
                "Succ Building": succ_info.get("Building", ""),
                "Succ Level": succ_info.get("Level", ""),
                "Succ Phase": succ_info.get("Phase", ""),
            })

    return pd.DataFrame(rows)[ordered_columns]


def _sanitize_wbs_id_token(value, fallback="X"):
    token = re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())
    return token or fallback


def _safe_duration(value, default=0.0):
    try:
        return float(value) if value not in (None, "") else float(default)
    except:
        return float(default)


def _phase_bucket(phase_value):
    text = str(phase_value or "").strip().upper()
    if text == "REBAR": return ("Rebar", "R", 1)
    if text in {"FORMWORK", "DECKING", "EP+FWK", "EMBEDMENTS"}: return ("EP+FWK", "F", 2)
    if text == "CONCRETE": return ("Concrete", "C", 3)
    if text == "INSTALLATION": return ("Precast Installation", "I", 4)
    if text == "STITCHING": return ("Stitching", "S", 5)
    pretty = str(phase_value or "").strip().title() or "Other"
    return (pretty, "O", 9)


def _build_synchro_wbs_table(df_activities):
    ordered_columns = [
        "No.", "Node Type", "WBS Line ID", "Parent ID", "Activity ID", "Activity Name",
        "Original Duration", "Building", "Level", "WBS", "Phase Group", "Task ID",
        "Task Name", "PlanningProposedPourID", "Include"
    ]

    if df_activities is None or df_activities.empty:
        return pd.DataFrame(columns=ordered_columns)

    df = df_activities.copy()
    for col in ["Building", "Level", "WBS", "Phase", "Days", "Task ID", "Task Name", "PlanningProposedPourID", "Include"]:
        if col not in df.columns:
            df[col] = ""

    df["_Duration"] = df["Days"].apply(_safe_duration)
    df["_LevelSortKey"] = df["Level"].apply(_extract_level_sort_key)
    phase_meta = df["Phase"].apply(_phase_bucket)
    df["_PhaseGroup"] = phase_meta.apply(lambda x: x[0])
    df["_PhaseCode"] = phase_meta.apply(lambda x: x[1])
    df["_PhaseSort"] = phase_meta.apply(lambda x: x[2])

    df = df.sort_values(
        by=["Building", "_LevelSortKey", "Level", "WBS", "_PhaseSort", "_PhaseGroup", "PlanningProposedPourID", "Task ID"],
        ascending=[True, True, True, True, True, True, True, True]
    ).reset_index(drop=True)

    rows = []

    def _append_row(node_type, wbs_line_id, parent_id, activity_id, activity_name, duration,
                    building="", level="", wbs_name="", phase_group="", task_id="", task_name="",
                    proposed_id="", include=""):
        rows.append({
            "No.": len(rows) + 1,
            "Node Type": node_type,
            "WBS Line ID": wbs_line_id,
            "Parent ID": parent_id,
            "Activity ID": activity_id,
            "Activity Name": activity_name,
            "Original Duration": round(float(duration), 3),
            "Building": building,
            "Level": level,
            "WBS": wbs_name,
            "Phase Group": phase_group,
            "Task ID": task_id,
            "Task Name": task_name,
            "PlanningProposedPourID": str(proposed_id or "").zfill(4) if str(proposed_id or "").strip() else "",
            "Include": include,
        })

    for building, df_building in df.groupby("Building", sort=False):
        building_text = str(building or "UNKNOWN")
        building_token = _sanitize_wbs_id_token(building_text, fallback="BLD")
        root_id = f"WBS-{building_token}-CON"
        root_duration = df_building["_Duration"].sum()
        _append_row("WBS_ROOT", root_id, "", root_id, f"{building_text} - Construction", root_duration, building=building_text)

        for level, df_level in df_building.groupby("Level", sort=False):
            level_text = str(level or "Unknown")
            level_token = _sanitize_wbs_id_token(level_text, fallback="LVL")
            level_id = f"{root_id}-{level_token}"
            level_duration = df_level["_Duration"].sum()
            _append_row("WBS_LEVEL", level_id, root_id, level_id, f"{building_text} - {level_text}", level_duration,
                        building=building_text, level=level_text)

            for group_index, (wbs_name, df_wbs) in enumerate(df_level.groupby("WBS", sort=False), 1):
                wbs_text = str(wbs_name or f"{building_text} - {level_text}")
                wbs_id = f"{level_id}-G{group_index:02d}"
                wbs_duration = df_wbs["_Duration"].sum()
                _append_row("WBS_GROUP", wbs_id, level_id, level_id, wbs_text, wbs_duration,
                            building=building_text, level=level_text, wbs_name=wbs_text)

                for phase_group, df_phase in df_wbs.groupby("_PhaseGroup", sort=False):
                    phase_code = _sanitize_wbs_id_token(df_phase["_PhaseCode"].iloc[0], fallback="P")
                    phase_id = f"{wbs_id}-{phase_code}"
                    phase_duration = df_phase["_Duration"].sum()
                    _append_row("WBS_PHASE", phase_id, wbs_id, phase_id, f"{wbs_text} - {phase_group}", phase_duration,
                                building=building_text, level=level_text, wbs_name=wbs_text, phase_group=phase_group)

                    for _, task in df_phase.iterrows():
                        task_id = str(task.get("Task ID", "") or "").strip()
                        task_name = str(task.get("Task Name", "") or "").strip() or f"{wbs_text} - {phase_group}"
                        _append_row("TASK", "", phase_id, task_id, task_name, task.get("_Duration", 0.0),
                                    building=building_text, level=level_text, wbs_name=wbs_text,
                                    phase_group=phase_group, task_id=task_id, task_name=task_name,
                                    proposed_id=task.get("PlanningProposedPourID", ""),
                                    include=task.get("Include", ""))

    if not rows:
        return pd.DataFrame(columns=ordered_columns)
    return pd.DataFrame(rows)[ordered_columns]


_SEQUENCE_ORDER_COLUMN_ORDER = [
    "Index", "Pour ID", "Level", "General Level 3", "Pour Type", "Concrete Stage", "Pour Volume m3", "Height",
    "Logic Group 1", "Logic Group 2", "Logic Group 3",
    ALT_POUR_ID_PROP, "PlanningProposedPourID", "GUID",
]


def _coerce_index_column(df):
    """Force the "Index" column to real numbers (pandas nullable Int64) so
    Excel stores it as numeric rather than text — blanks stay blank. Used
    for both the regenerated "SequenceOrder Template" and the round-tripped
    "SequenceOrder" sheet (which is read back with dtype=str, so its Index
    would otherwise be re-written as text)."""
    if df is not None and "Index" in df.columns:
        df = df.copy()
        df["Index"] = pd.to_numeric(df["Index"], errors="coerce").astype("Int64")
    return df


def _reorder_sequence_order_columns(df):
    """Puts the known SequenceOrder columns in their usual order; any column
    the user added manually that isn't part of that known set is kept,
    appended at the end, rather than silently dropped."""
    known = [c for c in _SEQUENCE_ORDER_COLUMN_ORDER if c in df.columns]
    extra = [c for c in df.columns if c not in _SEQUENCE_ORDER_COLUMN_ORDER]
    return df[known + extra]


_FROM_MODEL_COLUMN_NAMES = {"From Model", "#", "Object Name", "CWA Pour.Pour ID Number", "IFC.GlobalId"}


def _strip_from_model_columns(df):
    """Drops the now-retired "From Model" column and its "tbl_fromModel"
    lookup block (# / Object Name / CWA Pour.Pour ID Number / IFC.GlobalId,
    plus blank spacer columns) that older runs wrote onto the SequenceOrder
    sheet — so an Overrides file created before this feature was removed
    gets cleaned up on its next regeneration instead of carrying the stale
    block forward (or, worse, having it re-appended further right each run)."""
    if df is None:
        return df
    keep = [
        col for col in df.columns
        if re.sub(r"\.\d+$", "", str(col)) not in _FROM_MODEL_COLUMN_NAMES
        and not str(col).startswith("Unnamed:")
    ]
    return df[keep]


# ========================== MAIN EXPORT FUNCTION ==========================

def export_activities_to_excel(activity_list, excel_path: str, clash_results=None, records=None,
                               overrides=None, numbering_mode="sequential", task_predecessors=None,
                               task_start_dates=None, skipped_edges_by_source=None):
    if records is None:
        records = []
    if clash_results is None:
        clash_results = []
    if overrides is None:
        overrides = {}

    pour_lookup = {r["element"].GlobalId: r for r in records 
                   if "element" in r and hasattr(r["element"], "GlobalId")}

    rows = []
    for i, act in enumerate(activity_list, 1):
        r = act["rec"]
        phase = act["phase"]
        building = r.get("building", "UNKNOWN")
        conc_stage = str(r.get("conc_stage", "")).strip().upper()

        if conc_stage == "PRECAST":
            short_phase = {"Installation": "Precast Installation", "Stitching": "Stitching"}.get(phase, phase)
        else:
            short_phase = {"Rebar": "Rebar", "Formwork": "EP+FWK", "Concrete": "Concrete"}.get(phase, phase)

        lg1 = r.get("lg1", "??")
        pour_id = r.get("pour_id", "??")
        level_code = r.get("level_code", "??").upper()

        task_name = act.get("task_name", f"{building} - {level_code} - {lg1} {pour_id} - {short_phase}")

        row = {
            "No.": i,
            "PlanningProposedPourID": r.get("planning_proposed_pour_id", ""),
            "WBS": r.get("concrete_wbs_name", ""),
            "Phase": phase,
            "Task ID": act["task_id"],
            "Task Name": task_name,
            "Full Pour ID": r.get("full_pour_id", ""),
            ALT_POUR_ID_PROP: r.get("alt_pour_id", r.get("original_pour_id", "")),
            "Predecessors": str(act.get("predecessor") or ""),
            "Days": act.get("duration_days", 0),
            "Qty": round(r.get("volume", 0), 3) if phase in ("Concrete", "Stitching") else
                 round(r.get("rebar_weight_tonnes", 0), 3) if phase == "Rebar" else "",
            "Quantity Source": r.get("quantity_source", "Unknown"),
            "Level": level_code,
            "Ext": "YES" if r.get("is_external", False) else "NO",
            "Stage": r.get("conc_stage", ""),
            "Zone": r.get("building_zone", ""),
            "Type": r.get("pour_type", "").title(),
            "Logic Group 1": lg1,
            "Logic Group 2": r.get("lg2", ""),
            "Logic Group 3": r.get("lg3", ""),
            "Max Z": r.get("max_elevation", ""),
            "Min Z": r.get("min_elevation", ""),
            "Include": r.get("include", "NO"),
            "Pour ID": pour_id,
            "Resource": act.get("resource", ""),
            "Building": building,
            "GUID": r["element"].GlobalId if "element" in r else "MISSING_GUID",
        }
        rows.append(row)

    df_activities = pd.DataFrame(rows)

    if "PlanningProposedPourID" in df_activities.columns:
        df_activities["PlanningProposedPourID"] = df_activities["PlanningProposedPourID"].astype(str).str.strip().str.zfill(4)

    if "Level" in df_activities.columns:
        df_activities["_LevelSortKey"] = df_activities["Level"].apply(_extract_level_sort_key)
    else:
        df_activities["_LevelSortKey"] = 999999

    # Connected Pours
    clash_dict = defaultdict(set)
    for c in clash_results:
        a_guid = c.get("Pour A GUID") or c.get("a_global_id", "")
        b_guid = c.get("Pour B GUID") or c.get("b_global_id", "")
        if a_guid and b_guid:
            clash_dict[a_guid].add(b_guid)
            clash_dict[b_guid].add(a_guid)

    if "GUID" in df_activities.columns:
        df_activities["Connected Pours"] = df_activities["GUID"].apply(
            lambda g: len(clash_dict.get(g, set())) if g != "MISSING_GUID" else 0
        )

    # Sorting
    if numbering_mode == "stage_priority":
        df_activities = df_activities.sort_values(
            by=["Include", "Building", "_LevelSortKey", "Ext", "Type", "Stage", "Logic Group 1", "Logic Group 2", "Phase", "PlanningProposedPourID"],
            ascending=[False, True, True, False, True, True, False, False, True, True]
        ).reset_index(drop=True)
    else:
        df_activities = df_activities.sort_values(
            by=["Include", "Building", "_LevelSortKey", "Ext", "Type", "Stage", "Logic Group 1", "Logic Group 2", "Full Pour ID", "Phase", "PlanningProposedPourID"],
            ascending=[False, True, True, False, True, True, False, False, True, True, True]
        ).reset_index(drop=True)

    if "_LevelSortKey" in df_activities.columns:
        df_activities = df_activities.drop(columns=["_LevelSortKey"])

    df_activities["No."] = range(1, len(df_activities) + 1)

    # Build supporting dataframes
    _tn0 = _time.perf_counter()
    df_numbering_summary = _build_numbering_summary_from_records(records, numbering_mode=numbering_mode)
    df_numbering = _build_numbering_table_from_records(records, overrides=overrides, include_override_column=True, numbering_mode=numbering_mode)
    print(f"  [TIMING] excel: numbering tables: {_time.perf_counter() - _tn0:.2f}s")
    _tn1 = _time.perf_counter()
    df_synchro_wbs = _build_synchro_wbs_table(df_activities)
    df_clash_relations = _build_clash_relations_from_records(clash_results, records)
    df_duplicate_pour_ids = _build_duplicate_pour_ids_table(records)
    df_sequence_order_template = _build_sequence_order_template(records, activity_list, task_predecessors, task_start_dates)
    df_sequence_order_conflicts = _build_sequence_order_conflicts_table(activity_list, skipped_edges_by_source)
    df_building_mapping = _build_building_mapping_table(records)
    print(f"  [TIMING] excel: other supporting tables: {_time.perf_counter() - _tn1:.2f}s")

    # ====================== WRITE MAIN EXCEL ======================
    _tn2 = _time.perf_counter()
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        df_activities.to_excel(writer, sheet_name="Activities", index=False)
        df_numbering_summary.to_excel(writer, sheet_name="Numbering Summary", index=False)
        df_numbering.to_excel(writer, sheet_name="Numbering", index=False)
        df_synchro_wbs.to_excel(writer, sheet_name="Synchro WBS", index=False)
        if not df_duplicate_pour_ids.empty:
            df_duplicate_pour_ids.to_excel(writer, sheet_name="Duplicate Pour IDs", index=False)
        if not df_building_mapping.empty:
            df_building_mapping.to_excel(writer, sheet_name="Building Mapping", index=False)

        # Group Clashes
        clash_rows = []
        for idx, c in enumerate(clash_results, 1):
            a_guid = c.get("Pour A GUID") or c.get("a_global_id", "")
            b_guid = c.get("Pour B GUID") or c.get("b_global_id", "")
            a = pour_lookup.get(a_guid, {})
            b = pour_lookup.get(b_guid, {})
            clash_rows.append({
                "No.": idx,
                "A:Include": "YES" if a.get("include", "NO") == "YES" else "NO",
                "A:IsExternal": "YES" if a.get("is_external", False) else "NO",
                "A:PourType": a.get("pour_type", "").title(),
                "A:ConcreteStage": a.get("conc_stage", ""),
                "A:LogicGroup1": a.get("lg1", ""),
                "A:PourID": a.get("pour_id", ""),
                f"A:{ALT_POUR_ID_PROP}": a.get("alt_pour_id", a.get("original_pour_id", "")),
                "A:Min Z": a.get("min_elevation", ""),
                "A:Max Z": a.get("max_elevation", ""),
                "B:Include": "YES" if b.get("include", "NO") == "YES" else "NO",
                "B:IsExternal": "YES" if b.get("is_external", False) else "NO",
                "B:PourType": b.get("pour_type", "").title(),
                "B:ConcreteStage": b.get("conc_stage", ""),
                "B:LogicGroup1": b.get("lg1", ""),
                "B:PourID": b.get("pour_id", ""),
                f"B:{ALT_POUR_ID_PROP}": b.get("alt_pour_id", b.get("original_pour_id", "")),
                "B:Min Z": b.get("min_elevation", ""),
                "B:Max Z": b.get("max_elevation", ""),
                "Position (B rel. A)": classify_vertical_position(a, b) if a and b else "Undetermined",
                "Distance (mm)": c.get("Distance (mm)", ""),
                "Distance (m)": c.get("Distance (m)", ""),
                "Violation": c.get("Violation", "< 5mm clearance")
            })

        df_group_clashes = pd.DataFrame(clash_rows)
        if df_group_clashes.empty:
            df_group_clashes = pd.DataFrame(columns=["No.", "A:Include", "A:IsExternal", "A:PourType", "A:ConcreteStage",
                                                     "A:LogicGroup1", "A:PourID", f"A:{ALT_POUR_ID_PROP}", "A:Min Z",
                                                     "A:Max Z", "B:Include", "B:IsExternal", "B:PourType", "B:ConcreteStage",
                                                     "B:LogicGroup1", "B:PourID", f"B:{ALT_POUR_ID_PROP}", "B:Min Z",
                                                     "B:Max Z", "Position (B rel. A)", "Distance (mm)", "Distance (m)", "Violation"])
        df_group_clashes.to_excel(writer, sheet_name="Group Clashes", index=False)
        df_clash_relations.to_excel(writer, sheet_name="Clash Relations", index=False)

        # Formatting/tables applied to writer.book directly (the SAME
        # in-memory openpyxl Workbook every sheet above was just written
        # into) instead of closing this writer, then separately re-opening
        # the file from disk with load_workbook and saving it a second time
        # — that reopen+resave round trip used to cost ~9s on its own
        # (~5s load_workbook + ~3.5s wb.save, measured), on top of writing
        # the file in the first place, for a large multi-sheet workbook.
        main_widths = {
            "Activities": _fast_column_widths(df_activities),
            "Numbering Summary": _fast_column_widths(df_numbering_summary),
            "Numbering": _fast_column_widths(df_numbering),
            "Synchro WBS": _fast_column_widths(df_synchro_wbs),
            "Duplicate Pour IDs": _fast_column_widths(df_duplicate_pour_ids),
            "Building Mapping": _fast_column_widths(df_building_mapping),
            "Group Clashes": _fast_column_widths(df_group_clashes),
            "Clash Relations": _fast_column_widths(df_clash_relations),
        }
        _apply_sheet_formatting_and_tables(writer.book, main_widths)
    print(f"  [TIMING] excel: write+format main workbook: {_time.perf_counter() - _tn2:.2f}s")

    # ====================== REGENERATE OVERRIDES FILE ======================
    _tn3 = _time.perf_counter()
    overrides_path = Path(excel_path).parent / f"{Path(excel_path).stem.replace('_CWA_Planning', '')}_Overrides.xlsx"
    overrides_widths = {
        "SequenceOrder Template": _fast_column_widths(df_sequence_order_template),
        "SequenceOrder Conflicts": _fast_column_widths(df_sequence_order_conflicts),
    }

    is_first_generation = not overrides_path.exists()
    existing_sequence_order = None
    existing_relations = None
    existing_pour_id_order = None
    if overrides_path.exists():
        try:
            existing_sequence_order = pd.read_excel(str(overrides_path), sheet_name="SequenceOrder", dtype=str)
            existing_sequence_order = _strip_from_model_columns(existing_sequence_order)
        except:
            existing_sequence_order = None
        try:
            existing_relations = pd.read_excel(str(overrides_path), sheet_name="Relations", dtype=str)
        except:
            existing_relations = None
        try:
            existing_pour_id_order = pd.read_excel(str(overrides_path), sheet_name="PourIDOrder", dtype=str)
        except:
            existing_pour_id_order = None
        overrides_widths["SequenceOrder"] = _fast_column_widths(existing_sequence_order)
        overrides_widths["Relations"] = _fast_column_widths(existing_relations)
        overrides_widths["PourIDOrder"] = _fast_column_widths(existing_pour_id_order)

    override_rows = []
    for act in activity_list:
        rec = act["rec"]
        guid = rec["element"].GlobalId if "element" in rec and hasattr(rec["element"], "GlobalId") else "MISSING_GUID"
        phase_name = act["phase"]
        phase_key = f"{guid}_{phase_name}"
        ov = overrides.get(phase_key) or overrides.get(guid, {})

        row = {
            "No.": "",
            "GUID": guid,
            "Building": rec.get("building", "UNKNOWN"),
            "PlanningProposedPourID": rec.get("planning_proposed_pour_id", ""),
            "Level": rec.get("level_code", ""),
            "Pour ID": rec.get("pour_id", ""),
            ALT_POUR_ID_PROP: rec.get("alt_pour_id", rec.get("original_pour_id", "")),
            "FullPourID": rec.get("full_pour_id", ""),
            "Type": rec.get("pour_type", "").title(),
            "Ext": "YES" if rec.get("is_external", False) else "NO",
            "Zone": rec.get("building_zone", ""),
            "Stage": rec.get("conc_stage", ""),
            "Logic Group 1": rec.get("lg1", ""),
            "Logic Group 2": rec.get("lg2", ""),
            "Logic Group 3": rec.get("lg3", ""),
            "Max Z": rec.get("max_elevation", ""),
            "Min Z": rec.get("min_elevation", ""),
            "Include": rec.get("include", "NO"),
            "WBS": rec.get("concrete_wbs_name", ""),
            "Phase": phase_name,
            "TaskID_Override": ov.get("TaskID", ""),
            "TaskName_Override": ov.get("TaskName", ""),
            "Note": "",
            "LastGeneratedTaskID": act.get("task_id", ""),
            "LastGeneratedTaskName": act.get("task_name", ""),
        }
        override_rows.append(row)

    if override_rows:
        df_overrides = pd.DataFrame(override_rows)

        if "PlanningProposedPourID" in df_overrides.columns:
            df_overrides["PlanningProposedPourID"] = df_overrides["PlanningProposedPourID"].apply(_normalize_proposed_id_text)

        df_overrides = df_overrides.sort_values(
            by=["Building", "PlanningProposedPourID", "Level", "Ext", "Logic Group 1", "Logic Group 2", "Pour ID", "Phase"],
            ascending=[True, True, True, False, True, True, True, True]
        ).reset_index(drop=True)
        df_overrides["No."] = range(1, len(df_overrides) + 1)

        df_numbering_overrides = _build_numbering_table_from_records(
            records, overrides=overrides, include_override_column=True, numbering_mode=numbering_mode
        )
        overrides_widths["Overrides"] = _fast_column_widths(df_overrides)
        overrides_widths["Numbering"] = _fast_column_widths(df_numbering_overrides)

        with pd.ExcelWriter(str(overrides_path), engine='openpyxl') as writer:
            # SequenceOrder written first so it lands as the leftmost tab —
            # it's the sheet actually worked in day to day; Overrides/
            # Numbering/etc. are more reference/output than daily input.
            if existing_sequence_order is not None and not existing_sequence_order.empty:
                existing_sequence_order = _reorder_sequence_order_columns(existing_sequence_order)
                existing_sequence_order = _coerce_index_column(existing_sequence_order)
                existing_sequence_order.to_excel(writer, sheet_name="SequenceOrder", index=False)
            elif is_first_generation:
                # No Overrides file existed for this model yet — create the
                # SequenceOrder sheet/table scaffold (headers only, no rows)
                # so the user has tbl_SequenceOrder ready to paste
                # "SequenceOrder Template" rows into.
                df_seq_scaffold = pd.DataFrame(columns=_SEQUENCE_ORDER_COLUMN_ORDER)
                df_seq_scaffold.to_excel(writer, sheet_name="SequenceOrder", index=False)

            df_overrides.to_excel(writer, sheet_name="Overrides", index=False)
            df_numbering_overrides.to_excel(writer, sheet_name="Numbering", index=False)

            if existing_relations is not None and not existing_relations.empty:
                existing_relations.to_excel(writer, sheet_name="Relations", index=False)
            if existing_pour_id_order is not None and not existing_pour_id_order.empty:
                existing_pour_id_order.to_excel(writer, sheet_name="PourIDOrder", index=False)

            # Reference sheet, not live data — shows every current record in its
            # default generated order so the user has something to copy rows
            # from into the "SequenceOrder" sheet above. Lives here (Overrides)
            # rather than the main schedule report since it's an input aid for
            # the very sheet this file is meant to hold, not an export/report.
            df_sequence_order_template.to_excel(writer, sheet_name="SequenceOrder Template", index=False)

            # Every SequenceOrder link that had to be dropped this run because
            # enforcing it would create a logic cycle against the automatic
            # rules (level-gate/clash/Heavy Precast) — regenerated fresh every
            # run, right next to the sheet it's diagnosing. Empty when nothing
            # conflicted.
            df_sequence_order_conflicts.to_excel(writer, sheet_name="SequenceOrder Conflicts", index=False)

            # Formatted in place (writer.book) for the same reason as the main
            # workbook above — no separate load_workbook/save round trip.
            _apply_sheet_formatting_and_tables(writer.book, overrides_widths)

        if not df_sequence_order_conflicts.empty:
            print(f"  ⚠️ {len(df_sequence_order_conflicts)} SequenceOrder conflict(s) — see 'SequenceOrder Conflicts' sheet in {overrides_path.name}")

        print(f"Overrides file updated: {overrides_path.name} ({len(override_rows)} rows)")
    print(f"  [TIMING] excel: write+format overrides file: {_time.perf_counter() - _tn3:.2f}s")

    print(f"Excel report generated: {excel_path}")


def merge_excel_files(source_excels, master_excel):
    print(f"\n{'='*60}\nMERGING EXCEL FILES INTO MASTER\n{'='*60}")

    if not source_excels:
        print("No Excel files to merge")
        return

    all_activities = []
    all_numbering = []
    all_clashes = []
    all_relations = []
    all_duplicate = []
    all_building_mapping = []

    for excel_path in source_excels:
        if not excel_path.exists():
            print(f"Warning: {excel_path.name} not found, skipping...")
            continue
        print(f"Reading: {excel_path.name}")

        try:
            df_act = pd.read_excel(excel_path, sheet_name="Activities")
            df_act["Source File"] = excel_path.stem
            all_activities.append(df_act)

            try:
                df_num = pd.read_excel(excel_path, sheet_name="Numbering", dtype=str)
                df_num["Source File"] = excel_path.stem
                all_numbering.append(df_num)
            except:
                pass

            try:
                df_clash = pd.read_excel(excel_path, sheet_name="Group Clashes")
                df_clash["Source File"] = excel_path.stem
                all_clashes.append(df_clash)
            except:
                pass

            try:
                df_rel = pd.read_excel(excel_path, sheet_name="Clash Relations")
                df_rel["Source File"] = excel_path.stem
                all_relations.append(df_rel)
            except:
                pass

            try:
                df_dup = pd.read_excel(excel_path, sheet_name="Duplicate Pour IDs")
                df_dup["Source File"] = excel_path.stem
                all_duplicate.append(df_dup)
            except:
                pass

            try:
                df_bmap = pd.read_excel(excel_path, sheet_name="Building Mapping")
                df_bmap["Source File"] = excel_path.stem
                all_building_mapping.append(df_bmap)
            except:
                pass
        except Exception as e:
            print(f"  Error reading {excel_path.name}: {e}")
            continue

    if not all_activities:
        print("No activities found to merge")
        return

    df_master_activities = pd.concat(all_activities, ignore_index=True)
    if "PlanningProposedPourID" in df_master_activities.columns:
        df_master_activities["PlanningProposedPourID"] = df_master_activities["PlanningProposedPourID"].astype(str).str.strip().str.zfill(4)

    if "Level" in df_master_activities.columns:
        df_master_activities["_LevelSortKey"] = df_master_activities["Level"].apply(_extract_level_sort_key)

    df_master_activities = df_master_activities.sort_values(
        by=["Include", "Building", "PlanningProposedPourID", "_LevelSortKey", "Ext", "Type", "Stage", "Logic Group 1", "Logic Group 2", "Full Pour ID", "Phase"],
        ascending=[False, True, True, True, False, True, True, False, False, True, True]
    ).reset_index(drop=True)

    if "_LevelSortKey" in df_master_activities.columns:
        df_master_activities = df_master_activities.drop(columns=["_LevelSortKey"])

    df_master_activities["No."] = range(1, len(df_master_activities) + 1)

    df_master_clashes = pd.concat(all_clashes, ignore_index=True) if all_clashes else pd.DataFrame()
    df_master_relations = pd.concat(all_relations, ignore_index=True) if all_relations else pd.DataFrame()
    df_master_numbering = pd.concat(all_numbering, ignore_index=True) if all_numbering else pd.DataFrame()
    df_master_dup = pd.concat(all_duplicate, ignore_index=True) if all_duplicate else pd.DataFrame()
    if df_master_dup.empty:
        df_master_dup = pd.DataFrame(columns=[
            "No.", "Former Pour ID", "Objects with same ID", "GUID",
            "Proposed Pour ID", "Building", "Level", "Pour Type", "Source File"
        ])

    df_master_numbering_summary = _build_numbering_summary_from_activities(df_master_activities) if 'df_master_activities' in locals() else pd.DataFrame()
    df_master_synchro_wbs = _build_synchro_wbs_table(df_master_activities)
    df_master_building_mapping = pd.concat(all_building_mapping, ignore_index=True) if all_building_mapping else pd.DataFrame()

    with pd.ExcelWriter(str(master_excel), engine='openpyxl') as writer:
        df_master_activities.to_excel(writer, sheet_name="Activities", index=False)
        df_master_numbering_summary.to_excel(writer, sheet_name="Numbering Summary", index=False)
        df_master_numbering.to_excel(writer, sheet_name="Numbering", index=False)
        df_master_synchro_wbs.to_excel(writer, sheet_name="Synchro WBS", index=False)
        df_master_clashes.to_excel(writer, sheet_name="Group Clashes", index=False)
        df_master_relations.to_excel(writer, sheet_name="Clash Relations", index=False)
        df_master_dup.to_excel(writer, sheet_name="Duplicate Pour IDs", index=False)
        if not df_master_building_mapping.empty:
            df_master_building_mapping.to_excel(writer, sheet_name="Building Mapping", index=False)

    print(f"\nMASTER EXCEL CREATED: {master_excel.name}")
    print(f"Total Activities: {len(df_master_activities)}")
    print(f"Total Files merged: {len(source_excels)}")